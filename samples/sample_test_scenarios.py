"""
Script to create a sample Word document with test scenarios.
Run this script to generate a demo Word document for testing.
"""

from docx import Document
from docx.shared import Inches
from pathlib import Path


def create_sample_document():
    """Create a sample Word document with test scenarios."""
    
    # Create a new document
    document = Document()
    
    # Add title
    title = document.add_heading('Asset Management - Test Scenarios', 0)
    
    # Add introduction
    intro = document.add_paragraph(
        'This document contains functional test scenarios for the Asset Management system. '
        'Each test scenario includes a unique identifier, description, preconditions, and expected results.'
    )
    
    # Add section heading
    section = document.add_heading('Functional Test Scenarios', level=1)
    
    # Create table with test scenarios
    table = document.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    
    # Add header row
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Scenario ID'
    hdr_cells[1].text = 'Description'
    hdr_cells[2].text = 'Preconditions'
    hdr_cells[3].text = 'Expected Results'
    
    # Sample test scenarios data
    test_scenarios = [
        {
            'id': 'AM-001',
            'description': 'Create New Asset Record',
            'preconditions': 'User is logged in with asset creation permissions. Asset creation form is accessible.',
            'expected_results': 'New asset record is created successfully with unique asset ID. Asset appears in asset registry with all provided details.'
        },
        {
            'id': 'AM-002', 
            'description': 'Search Asset by Asset ID',
            'preconditions': 'Asset records exist in the system. User has search permissions.',
            'expected_results': 'System returns the correct asset record matching the provided Asset ID. Asset details are displayed accurately.'
        },
        {
            'id': 'AM-003',
            'description': 'Update Asset Information',
            'preconditions': 'Asset record exists in system. User has update permissions for the asset.',
            'expected_results': 'Asset information is updated successfully. Updated timestamp is recorded. Change history is maintained.'
        },
        {
            'id': 'AM-004',
            'description': 'Generate Asset Report',
            'preconditions': 'Asset data exists in system. User has reporting permissions.',
            'expected_results': 'Asset report is generated with current data. Report includes all requested fields and filters. Report can be exported successfully.'
        },
        {
            'id': 'AM-005',
            'description': 'Delete Asset Record',
            'preconditions': 'Asset record exists and is not referenced by other records. User has deletion permissions.',
            'expected_results': 'Asset record is marked as deleted or moved to archive. Asset no longer appears in active asset searches.'
        },
        {
            'id': 'AM-006',
            'description': 'Assign Asset to Employee',
            'preconditions': 'Asset exists and is available for assignment. Employee record exists in system.',
            'expected_results': 'Asset is successfully assigned to employee. Assignment record is created with timestamp. Asset status updated to "Assigned".'
        },
        {
            'id': 'AM-007',
            'description': 'Track Asset Location History',
            'preconditions': 'Asset has been moved between different locations. Location tracking is enabled.',
            'expected_results': 'System displays complete location history with timestamps. Current location is accurately reflected.'
        },
        {
            'id': 'AM-008',
            'description': 'Set Asset Maintenance Schedule',
            'preconditions': 'Asset record exists. User has maintenance scheduling permissions.',
            'expected_results': 'Maintenance schedule is created and associated with asset. Notifications are set for upcoming maintenance dates.'
        }
    ]
    
    # Add test scenario rows
    for scenario in test_scenarios:
        row_cells = table.add_row().cells
        row_cells[0].text = scenario['id']
        row_cells[1].text = scenario['description']
        row_cells[2].text = scenario['preconditions']
        row_cells[3].text = scenario['expected_results']
    
    # Add notes section
    notes = document.add_heading('Notes', level=2)
    notes_para = document.add_paragraph(
        '• All test scenarios should be executed in a test environment\n'
        '• Ensure proper test data setup before execution\n'
        '• Document any deviations from expected results\n'
        '• Test scenarios can be automated using the Xray Test Automation tool'
    )
    
    # Save the document
    output_path = Path(__file__).parent / 'sample_test_scenarios.docx'
    document.save(str(output_path))
    
    print(f"Sample document created: {output_path}")
    return output_path


if __name__ == "__main__":
    create_sample_document()